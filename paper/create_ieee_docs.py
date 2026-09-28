import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = '/Users/anilkumar/Downloads/ORCHID/paper'
IMG_DIR = os.path.join(BASE_DIR, 'extracted_images')

def set_cell_margins(cell, top=50, bottom=50, left=100, right=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for border_name, border_style in borders.items():
        if border_style:
            b_el = OxmlElement(f'w:{border_name}')
            b_el.set(qn('w:val'), border_style.get('val', 'single'))
            b_el.set(qn('w:sz'), str(border_style.get('sz', 4)))
            b_el.set(qn('w:space'), '0')
            b_el.set(qn('w:color'), border_style.get('color', '000000'))
            tcBorders.append(b_el)
        else:
            b_el = OxmlElement(f'w:{border_name}')
            b_el.set(qn('w:val'), 'none')
            tcBorders.append(b_el)
    tcPr.append(tcBorders)

def set_two_columns(section):
    sectPr = section._sectPr
    cols = sectPr.xpath('./w:cols')
    if cols:
        col = cols[0]
    else:
        col = OxmlElement('w:cols')
        sectPr.append(col)
    col.set(qn('w:num'), '2')
    col.set(qn('w:space'), '360')  # 0.25 inch space (360 dxa)

def build_docx():
    doc = Document()
    
    # Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    
    # Section 1: Title, Authors, Page Margins
    sec1 = doc.sections[0]
    sec1.page_width = Inches(8.5)
    sec1.page_height = Inches(11.0)
    sec1.top_margin = Inches(0.75)
    sec1.bottom_margin = Inches(1.0)
    sec1.left_margin = Inches(0.63)
    sec1.right_margin = Inches(0.63)
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(10)
    p_title.paragraph_format.line_spacing = 1.1
    r_title = p_title.add_run("Conditional StyleGAN2-ADA for Synthetic Oral Histopathology Image Generation: A Pure-Synthetic Downstream Classification Evaluation on the ORCHID Dataset")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(20)
    r_title.bold = True
    
    # Authors
    p_author = doc.add_paragraph()
    p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_author.paragraph_format.space_before = Pt(0)
    p_author.paragraph_format.space_after = Pt(16)
    r_a1 = p_author.add_run("Chakkilala Anil Kumar        Malladi Nagasri        Bandaru Keerthana\n")
    r_a1.font.name = 'Times New Roman'
    r_a1.font.size = Pt(11)
    r_a1.bold = True
    
    r_affil = p_author.add_run("Department of Computer Science and Engineering\nORCHID Oral Histopathology Benchmark Study")
    r_affil.font.name = 'Times New Roman'
    r_affil.font.size = Pt(9)
    r_affil.italic = True
    
    # Section 2: Two-column body
    sec2 = doc.add_section(WD_SECTION_START.CONTINUOUS)
    sec2.page_width = Inches(8.5)
    sec2.page_height = Inches(11.0)
    sec2.top_margin = Inches(0.75)
    sec2.bottom_margin = Inches(1.0)
    sec2.left_margin = Inches(0.63)
    sec2.right_margin = Inches(0.63)
    set_two_columns(sec2)
    
    # Helpers
    def add_abstract(text, keywords=None):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.first_line_indent = Inches(0.14)
        
        r_lbl = p.add_run("Abstract—")
        r_lbl.bold = True
        r_lbl.italic = True
        r_lbl.font.size = Pt(9)
        
        r_txt = p.add_run(text)
        r_txt.bold = True
        r_txt.font.size = Pt(9)
        
        if keywords:
            pk = doc.add_paragraph()
            pk.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pk.paragraph_format.space_before = Pt(0)
            pk.paragraph_format.space_after = Pt(12)
            pk.paragraph_format.line_spacing = 1.0
            pk.paragraph_format.first_line_indent = Inches(0.14)
            
            rk_lbl = pk.add_run("Index Terms—")
            rk_lbl.bold = True
            rk_lbl.italic = True
            rk_lbl.font.size = Pt(9)
            
            rk_txt = pk.add_run(keywords)
            rk_txt.italic = True
            rk_txt.font.size = Pt(9)

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text.upper())
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.bold = True

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.italic = True
        r.bold = False

    def add_p(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.first_line_indent = Inches(0.14)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)

    def add_equation(eq_text, eq_num):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(eq_text + f"                      ({eq_num})")
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.italic = True

    def add_figure(img_path, caption_text, width=Inches(3.35)):
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(6)
            p_img.paragraph_format.space_after = Pt(2)
            p_img.paragraph_format.keep_with_next = True
            run_img = p_img.add_run()
            run_img.add_picture(img_path, width=width)
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(8)
            
            # Match "Fig. X."
            m = re.match(r'^(Fig\.\s*\d+\.)\s*(.*)$', caption_text)
            if m:
                rf = p_cap.add_run(m.group(1) + " ")
                rf.font.name = 'Times New Roman'
                rf.font.size = Pt(8)
                rf.bold = True
                
                rc = p_cap.add_run(m.group(2))
                rc.font.name = 'Times New Roman'
                rc.font.size = Pt(8)
            else:
                rc = p_cap.add_run(caption_text)
                rc.font.name = 'Times New Roman'
                rc.font.size = Pt(8)

    # 1. Abstract
    abstract_text = (
        "Oral Squamous Cell Carcinoma (OSCC) and its premalignant precursor, Oral Submucous Fibrosis (OSF), "
        "pose a severe public health burden in South Asia; however, geographic data disparities and severe class imbalances "
        "significantly undermine the diagnostic accuracy of medical AI models. Aiming to overcome this scarcity, we explored "
        "the clinical viability of generating histopathology images using a conditional StyleGAN2-ADA model trained on the "
        "ORCHID dataset. In spite of achieving an impressive Fréchet Inception Distance (FID) of 6.98, comprehensive downstream "
        "testing exposed a substantial discrepancy between visual fidelity and downstream diagnostic utility. In synthetic-only "
        "training regimes, classifiers (DenseNet-121 and EfficientNet-B3) achieved below 10% accuracy when tested against "
        "authentic clinical biopsies, indicating a failure to encode fine-grained microscopic features. Subsequent data-mixing "
        "experiments revealed that while conservative synthetic supplementation effectively balances minority classes without "
        "compromising diagnostic accuracy, excessive synthetic saturation actively dilutes genuine biological signals, resulting "
        "in performance decline. Overall, this study highlights the clinical deception of purely visual realism; the results "
        "support using synthetic data conservatively as a complementary augmentation strategy rather than as an unconstrained "
        "replacement for real patient tissue."
    )
    keywords = "Oral Histopathology, StyleGAN2-ADA, ORCHID Dataset, Synthetic Data Generalization, Pure-Synthetic Stress Test, DenseNet-121, EfficientNet-B3."
    add_abstract(abstract_text, keywords)
    
    # 2. Section I
    add_heading_1("I. Introduction")
    add_p(
        "In South Asia, the disproportionate burden of cancer-related morbidity and mortality is driven significantly by oral "
        "cancer, particularly Oral Squamous Cell Carcinoma (OSCC) and its high-risk precursor state, Oral Submucous Fibrosis (OSF). "
        "The etiology of oral malignancies in South Asia is strongly linked to habit-related carcinogens, such as consumption of "
        "betel quid, gutka, and smokeless tobacco, consequently yielding distinct genomic, inflammatory, and structural phenotypes. "
        "Despite this immense disease burden, clinical translation in computational pathology is bottlenecked by profound geographic "
        "representation bias. Since more than half of global clinical AI datasets originate from Western or East Asian centres, neural "
        "networks frequently exploit non-clinical shortcuts instead of learning valid histological features. As a result, systems trained "
        "on foreign datasets routinely experience catastrophic domain generalization failures when deployed in Indian clinical environments."
    )
    add_p(
        "To mitigate this representation gap, the multi-center ORCHID dataset was established, compiling diverse H&E oral "
        "histopathology images across five diagnostic categories. Yet, consistent with clinical data collection, the ORCHID dataset "
        "contains moderate class imbalances. Consequently, this reduces the model sensitivity on minority, high-complexity classes, "
        "including normal oral mucosa and aggressive PDOSCC. To address this data bottleneck, generative modelling has emerged as a "
        "promising countermeasure. Generative Adversarial Networks (GANs), specifically StyleGAN2-ADA, offer a promising framework "
        "to synthesize high-resolution tissue patches and balance underrepresented clinical cohorts."
    )
    add_p(
        "However, an essential question persists in medical AI: Does visual and statistical fidelity reliably translate into clinical "
        "diagnostic utility? Standard metrics like Fréchet Inception Distance (FID) rely on general-purpose feature extractors rather "
        "than microscopic diagnostic features. Focusing on global color and broad textures, these metrics are easily misled by visually "
        "plausible synthetic tissue lacking essential cellular atypia and nuclear pleomorphism."
    )
    add_p(
        "In this work, we investigate the divergence between statistical generative realism and clinical diagnostic utility through a "
        "rigorous, multi-stage empirical framework. We initially optimized a class-conditional StyleGAN2-ADA model on the ORCHID dataset "
        "to generate a benchmark repository of high-fidelity synthetic patches. To transcend standard statistical metrics, we evaluate "
        "downstream diagnostic validity using DenseNet-121 and EfficientNet-B3 architectures across three distinct experiments. We first "
        "tested classifiers trained purely on 25,000 synthetic images, followed by a controlled mixing setup where minority classes were "
        "supplemented to match majority-class size. Finally, we performed scale-expansion experiments by padding each diagnostic class up "
        "to 5,000 images. We applied the same α-controlled mixing formulation to govern the integration of real and synthetic "
        "histopathology patches."
    )
    
    # 3. Section II
    add_heading_1("II. Related Work")
    add_p(
        "Deep convolutional architectures such as DenseNet-121 [1] and EfficientNet-B3 [2] serve as standard backbones for fine-grained "
        "medical image classification. Despite their efficacy, the translation of these networks to South Asian oral oncology has historically "
        "been constrained by severe regional data deficits [3]. The dataset consists of approx. 300,000 raw high-resolution H&E-stained "
        "patches at 1000× effective magnification, and the final curated benchmark encompasses 14,705 images classified into five "
        "diagnostic categories: normal, OSMF, WDOSCC, MDOSCC, and PDOSCC [4]. Reflecting routine clinical intake, the dataset displays a "
        "moderate class imbalance ratio of approximately 2.67:1 [4], [5]. While WDOSCC and MDOSCC heavily dominate the distribution, "
        "normal tissue and aggressive PDOSCC cases remain severely underrepresented."
    )
    
    # Table I: ORCHID Dataset Class Distribution
    p_t1_lbl = doc.add_paragraph()
    p_t1_lbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t1_lbl.paragraph_format.space_before = Pt(6)
    p_t1_lbl.paragraph_format.space_after = Pt(2)
    p_t1_lbl.paragraph_format.keep_with_next = True
    r_t1 = p_t1_lbl.add_run("TABLE I\nORCHID DATASET CLASS DISTRIBUTION")
    r_t1.font.name = 'Times New Roman'
    r_t1.font.size = Pt(8)
    r_t1.bold = True
    
    t1 = doc.add_table(rows=7, cols=2)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    t1.autofit = False
    
    t1_data = [
        ("Diagnostic Class", "Number of Real Images"),
        ("Normal Oral Mucosa", "1,502"),
        ("Oral Submucous Fibrosis (OSMF)", "3,015"),
        ("Well-Differentiated OSCC (WDOSCC)", "4,011"),
        ("Moderately-Differentiated OSCC (MDOSCC)", "3,880"),
        ("Poorly-Differentiated OSCC (PDOSCC)", "2,297"),
        ("Total Cohort", "14,705")
    ]
    
    border_thick = {'val': 'single', 'sz': 12, 'color': '000000'}
    border_thin = {'val': 'single', 'sz': 4, 'color': '000000'}
    
    for r_idx, (col1, col2) in enumerate(t1_data):
        row = t1.rows[r_idx]
        cell1, cell2 = row.cells[0], row.cells[1]
        cell1.width = Inches(2.0)
        cell2.width = Inches(1.3)
        
        p1 = cell1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p1.add_run(col1)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(8)
        
        p2 = cell2.paragraphs[0]
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(2)
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r2 = p2.add_run(col2)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(8)
        
        if r_idx == 0:
            r1.bold = True
            r2.bold = True
            set_cell_borders(cell1, top=border_thick, bottom=border_thin)
            set_cell_borders(cell2, top=border_thick, bottom=border_thin)
        elif r_idx == 6:
            r1.bold = True
            r2.bold = True
            set_cell_borders(cell1, top=border_thin, bottom=border_thick)
            set_cell_borders(cell2, top=border_thin, bottom=border_thick)
        else:
            set_cell_borders(cell1)
            set_cell_borders(cell2)
        set_cell_margins(cell1, top=40, bottom=40, left=60, right=60)
        set_cell_margins(cell2, top=40, bottom=40, left=60, right=60)
        
    p_t1_sp = doc.add_paragraph()
    p_t1_sp.paragraph_format.space_after = Pt(6)

    add_p(
        "Prior studies utilizing the ORCHID benchmark have highlighted this distribution skew as a critical bottleneck limiting "
        "downstream diagnostic sensitivity. Previous whole-slide image (WSI) studies emphasized the scarcity of poorly differentiated "
        "carcinomas, identifying PDOSCC as a decisive bottleneck for automated grading [6]. To mitigate these skews, prior efforts either "
        "discarded majority-class data samples to achieve uniform class representation [7] or restricted the scope to binary "
        "classification (Normal versus OSCC) using stratified sampling and class-weighted loss to encounter a 6.78:1 imbalance [8]."
    )
    add_p(
        "Recently, long-tailed classification methods frame PDOSCC as a vulnerable tail category, implementing complex "
        "distribution-aware loss reweighting and tail-specific confidence gating to prevent F1-score collapse [9]. While these "
        "computational interventions effectively facilitate optimization convergence, they functionally treat the symptom of data skew "
        "through loss re-weighting rather than resolving the core biological dearth of representative tissue phenotypes."
    )
    add_p(
        "Seeking to resolve these biological deficits at the data level, Generative Adversarial Networks (GANs) are widely deployed "
        "to synthesize high-resolution histopathological imagery [10]. StyleGAN2-ADA is particularly well suited for oral histology "
        "synthesis because its style-based generator decouples global architectural patterns from subtle cellular details, and the adaptive "
        "discriminator augmentation (ADA) mitigates discriminator overfitting under acute sample scarcity [11]. Prior work has combined "
        "StyleGAN2-ADA with transfer learning in data-constrained domains such as breast cancer [11] and thyroid tissue [12], verifying "
        "that the architecture synthesizes photorealistic textures even under acute sample constraints."
    )
    add_p(
        "Despite these advancements, a pervasive methodological limitation persists in how synthetic histopathology is evaluated, "
        "with prevailing literature depending on the Fréchet Inception Distance (FID) as the primary arbiter of generative quality [13]. "
        "Because FID operates on natural-image features, it captures macroscopic textures while remaining blind to the fine microscopic "
        "markers required for histological grading [14], [15]. Moreover, existing literature evaluates synthetic samples only when "
        "combined with real images [15], [16], rather than testing them in isolation. As a result, existing literature stresses the "
        "absence of systematic, multimodal stress tests that evaluate synthetic images in complete isolation. Without evaluating the "
        "synthetic data in complete isolation, it is impossible to confirm whether low FID values reflect biological validity or merely "
        "superficial visual realism."
    )
    
    # 4. Section III
    add_heading_1("III. Methodology: Generative Synthesis and GAN Training")
    add_p(
        "To counteract data scarcity and class skew in the ORCHID repository, we deployed StyleGAN2-ADA with Adaptive Discriminator "
        "Augmentation (ADA) to stabilize training. To seed the network with generalized histological feature distributions, the model "
        "was initialized using BreCaHAD breast cancer weights to impart foundational histopathological priors and baseline H&E staining "
        "features. Training was performed in a class-conditional setting on an NVIDIA GeForce RTX 5090 GPU, integrating horizontal flips "
        "as dataset-level augmentation while maintaining an ADA target probability threshold of 0.6 to ensure stable adversarial dynamics. "
        "Generative quality was tracked throughout 3,000 kimg, employing 50,000-sample Fréchet Inception Distance (FID-50k) evaluations "
        "to identify the lowest-distance checkpoint for downstream validation. From a transfer baseline of 131.19, the FID-50k declined "
        "to a minimum of 6.98 at the 2,800-kimg checkpoint."
    )
    
    add_figure(
        os.path.join(IMG_DIR, "p3_img25.png"),
        "Fig. 1. FID score progression during GAN training highlighting the global minimum of 6.98 at 2,800 kimg."
    )
    
    add_p(
        "To quantify the trade-off between cellular-level realism and phenotypic variation, late-stage checkpoints were benchmarked "
        "against manifold-based Precision and Recall metrics."
    )
    
    add_figure(
        os.path.join(IMG_DIR, "p3_img23.png"),
        "Fig. 2. Checkpoint evaluation balancing image fidelity (Precision) against dataset diversity (Recall) across late-stage training snapshots."
    )
    
    add_p(
        "Checkpoint 2,800 demonstrated the most favourable fidelity-diversity trade-off while preserving substantial biological diversity "
        "relative to earlier snapshots (Precision: 0.6825, Recall: 0.4332). Based on this optimal trade-off, the 2,800 kimg weights were "
        "preserved to generate a dedicated bank of 25,000 synthetic patches for all downstream stress tests and augmentation experiments."
    )
    
    add_figure(
        os.path.join(IMG_DIR, "p3_img26.jpeg"),
        "Fig. 3. Qualitative comparison of synthetic patches generated by StyleGAN2-ADA against authentic clinical biopsies across the five ORCHID categories (Normal, OSMF, WDOSCC, MDOSCC, PDOSCC)."
    )
    
    # 5. Section IV
    add_heading_1("IV. Downstream Classification Methodology")
    add_heading_2("A. The Pure-Synthetic Stress Test")
    add_p(
        "Standard generative metrics, most notably the Fréchet Inception Distance (FID), evaluate statistical visual fidelity, "
        "fundamentally decoupling perceptual plausibility from downstream clinical veracity. To evaluate whether the generated tiles "
        "retained authentic diagnostic utility, we implemented a pure-synthetic stress test using zero real images during training. "
        "Under this zero-real experimental protocol, downstream classifiers were trained exclusively on an artificially synthesized "
        "repository of 25,000 patches (5,000 samples per class across the five-tier benchmark) with zero exposure to authentic "
        "clinical tissue during optimization. The diagnostic generalizability of these zero-real representations was strictly evaluated "
        "on an independent, holdout test set composed entirely of real clinical patient biopsies. By isolating synthetic representations "
        "during training, this protocol ensures that classifier performance solely depends on the generator’s ability to encode genuine "
        "microscopic morphology."
    )
    
    add_heading_2("B. α Ablation Strategy")
    add_p(
        "To identify the precise empirical threshold where synthetic data benefits or impairs classification, we conducted systematic "
        "mixing experiments combining synthetic and authentic cohorts. To calibrate synthetic augmentation, we defined a mixing parameter, "
        "α ∈ {0.0, 0.25, 0.50, 0.75, 1.0}, representing the proportion of the minority data deficit filled by generated patches. The "
        "volume of synthetic patches allocated to class c, denoted as S_c, is defined as:"
    )
    
    add_equation("S_c = [α × (T_target - N_c)]", "1")
    
    add_p(
        "where α denotes the fraction of the class deficit filled using generated images, N_c is the number of real images originally "
        "available in class c, and T_target is the required target volume in the combination of real and synthetic data. For example, if a "
        "class contains N_c = 1,500 real images and the target is T_target = 5,000, the deficit is 3,500 images. At α = 0.50, 1,750 "
        "synthetic images are added, producing a composite class size of 3,250."
    )
    add_p(
        "In the controlled balancing sweep (Experiment 3), T_target was fixed to the sample frequency of the primary majority class, "
        "standardizing the upper ceiling of each class prior to synthetic interpolation. On the other hand, in our aggressive dataset "
        "expansion sweep (Experiment 4), T_target was scaled to an artificial target of 5,000 images per class."
    )
    
    add_heading_2("C. Network Architectures and Preprocessing")
    add_p(
        "To ensure our findings were independent of architectural bias, we benchmarked two contrasting architectures: DenseNet-121, "
        "which utilizes iterative feature concatenation, and EfficientNet-B3, which applies adaptive compound scaling. For both "
        "architectures, the native ImageNet classifier of each network was replaced with a custom head consisting of a dropout layer "
        "and a linear projection to the five ORCHID diagnostic classes."
    )
    add_p(
        "To address varying data configurations across phases, spatial resolutions, augmentations, and optimization objectives were "
        "systematically adjusted to match the analytical targets of each evaluation phase. To standardize training dynamics, regularization "
        "techniques were applied, including label smoothing (ε = 0.1), gradient clipping bounded by a maximum norm of 1.0, and PyTorch "
        "automatic mixed precision (AMP) to maintain robust numerical convergence. The distinct preprocessing routines and hyperparameter "
        "settings for each classification experiment are detailed in Table II."
    )
    
    # Table II: Hyperparameters
    p_t2_lbl = doc.add_paragraph()
    p_t2_lbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t2_lbl.paragraph_format.space_before = Pt(6)
    p_t2_lbl.paragraph_format.space_after = Pt(2)
    p_t2_lbl.paragraph_format.keep_with_next = True
    r_t2 = p_t2_lbl.add_run("TABLE II\nPREPROCESSING AND HYPERPARAMETER CONFIGURATIONS ACROSS EXPERIMENTAL PHASES")
    r_t2.font.name = 'Times New Roman'
    r_t2.font.size = Pt(8)
    r_t2.bold = True
    
    t2 = doc.add_table(rows=4, cols=4)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t2.autofit = False
    
    t2_data = [
        ("Experimental Phase", "Exp 2: Pure Synthetic Stress Test", "Exp 3: Controlled Balancing Sweeps", "Exp 4: Aggressive Target Expansion"),
        ("Image Resolution & Augmentations", "300×300 px.\nRandom H/V flips, rotations up to 15°.", "224×224 px.\nRandomResizedCrop(0.8, 1.0), H/V flips, 15° rot, 15% color jitter.", "300×300 px.\nH/V flips, 15° rotations, standard color jittering."),
        ("Optimization & Scheduler", "25 Epochs (AdamW).\nOneCycleLR (Backbone: 10⁻⁴, Head: 10⁻³).", "25 Epochs (AdamW).\nCosineAnnealingLR (Base: 3×10⁻⁴, 3-epoch warmup).", "20 Epochs (AdamW).\nCosineAnnealingLR (Base: 10⁻³ tapering to 10⁻⁵)."),
        ("Dropout & Weight Decay", "Dropout: p = 0.4\nWD: 10⁻²", "Dropout: p = 0.3\nWD: 10⁻⁴", "Dropout: p = 0.3\nWD: 10⁻⁴")
    ]
    
    for r_idx, row_items in enumerate(t2_data):
        row = t2.rows[r_idx]
        for c_idx, cell_text in enumerate(row_items):
            cell = row.cells[c_idx]
            cell.width = Inches(0.85 if c_idx == 0 else 0.86)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(cell_text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(7.5)
            if r_idx == 0:
                r.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_cell_borders(cell, top=border_thick, bottom=border_thin)
            elif r_idx == 3:
                set_cell_borders(cell, top=border_thin, bottom=border_thick)
            else:
                set_cell_borders(cell)
            set_cell_margins(cell, top=30, bottom=30, left=40, right=40)
            
    p_t2_sp = doc.add_paragraph()
    p_t2_sp.paragraph_format.space_after = Pt(6)

    # 6. Section V
    add_heading_1("V. Experimental Results & Downstream Benchmarks")
    add_heading_2("A. Generative Benchmarking and the Pure-Synthetic Generalization Gap")
    add_p(
        "The conditional StyleGAN2-ADA framework demonstrated robust convergence, reaching an optimal FID-50k of 6.98 at the "
        "2,800-kimg snapshot before reaching asymptotic stability. This checkpoint established an optimal trade-off between sample "
        "fidelity (Precision = 0.6825) against structural diversity (Recall = 0.4332). Nonetheless, the pure-synthetic stress test "
        "revealed a significant gap between superficial generative realism and biological ground truth. When trained solely on the "
        "25,000-sample synthetic bank, both DenseNet-121 and EfficientNet-B3 rapidly exceeded 99% training accuracy with smooth loss decay."
    )
    
    add_figure(
        os.path.join(IMG_DIR, "p4_img29.png"),
        "Fig. 4. Synthetic optimization: Training accuracy trajectories of DenseNet-121 and EfficientNet-B3 under the pure-synthetic protocol."
    )
    add_figure(
        os.path.join(IMG_DIR, "p4_img30.png"),
        "Fig. 5. Synthetic optimization: Training loss decay curves over 25 epochs."
    )
    
    add_p(
        "When tested on authentic patient biopsies, both models suffered severe performance collapse, indicating that features "
        "learned from the synthetic manifold failed entirely to generalize to genuine pathological tissue specimens. As optimization "
        "advanced, validation loss diverged sharply upward, resulting in a severe generalization gap when tested against real patient biopsies."
    )
    
    add_figure(
        os.path.join(IMG_DIR, "p4_img31.png"),
        "Fig. 6. Generalization gap in DenseNet-121 and EfficientNet-B3 illustrating severe divergence between training accuracy on synthetic data and validation performance on genuine tissue."
    )
    add_figure(
        os.path.join(IMG_DIR, "p4_img33.png"),
        "Fig. 7. Pure-synthetic stress test: Real-world clinical collapse showing diagnostic accuracy dropping below 10% on real patient biopsies."
    )
    
    add_p(
        "Real-world diagnostic performance collapsed drastically to 9.17% for DenseNet-121 and 9.36% for EfficientNet-B3, accompanied "
        "by single-digit macro F1-scores. This profound ~90% collapse highlights a fundamental limitation of standard generative "
        "metrics, i.e., these metrics fail to assess whether synthetic images preserve authentic clinical markers. While StyleGAN2-ADA "
        "captured macro-level H&E staining profiles and broad stromal textures, the model failed to encode the subtle cytological "
        "markers critical for reliable clinical grading."
    )
    
    add_heading_2("B. Bounded Regularization via Controlled Balancing")
    add_p(
        "In the controlled balancing protocol, synthetic patches were introduced selectively to offset minority deficits, capping each "
        "class at the real majority volume and increasing the dataset from 10,228 to 13,950. Under this bounded regime, synthetic additions "
        "didn’t distort authentic biological variance and served as an effective regularizer for the downstream classifiers."
    )
    
    add_figure(
        os.path.join(IMG_DIR, "p5_img38.png"),
        "Fig. 8. Experiment 3: Test accuracy vs. mixing parameter α under bounded balancing sweeps."
    )
    add_figure(
        os.path.join(IMG_DIR, "p5_img39.png"),
        "Fig. 9. Experiment 3: Macro F1-score heatmap across α sweeps for DenseNet-121 and EfficientNet-B3."
    )
    
    add_p(
        "The two backbones exhibited divergent performance trajectories, driven by the distinct structural mechanics of each "
        "architecture. DenseNet-121 was sensitive to synthetic noise due to its dense feature concatenation, achieving its performance "
        "ceiling on the purely authentic baseline (96.80% test accuracy). In contrast, EfficientNet-B3 leveraged its compound scaling "
        "and channel attenuation to suppress high-frequency generative artifacts while exploiting the underlying morphological variance "
        "of the synthesized patches. This lifted EfficientNet-B3’s performance from an authentic baseline test accuracy of 96.29% to a "
        "performance zenith of 96.42% at full minority class balance (α = 1.00). This demonstrates that when constrained to authentic class "
        "limits, synthetic augmentation effectively mitigates class imbalance without compromising diagnostic accuracy."
    )
    
    add_heading_2("C. The Synthetic Dilution Threshold")
    add_p(
        "Experiment 4 evaluated the boundaries of generative regularization by expanding every diagnostic class to an artificial ceiling "
        "of 5,000 images per class. At maximum expansion (α = 1.00), the dataset swelled to 25,000 images, shifting the data distribution "
        "such that synthetic samples accounted for roughly 59% of the total volume. Such extensive saturation caused severe feature "
        "dilution, obscuring critical sub-cellular diagnostic features beneath repetitive, synthetic-manifold regularities. By overwhelming "
        "authentic tissue samples, generative artifacts corrupted the models’ discriminative decision boundaries and degraded grading accuracy."
    )
    
    add_figure(
        os.path.join(IMG_DIR, "p5_img41.png"),
        "Fig. 10. Experiment 4: The synthetic dilution threshold showing performance degradation beyond critical synthetic saturation."
    )
    add_figure(
        os.path.join(IMG_DIR, "p5_img40.png"),
        "Fig. 11. Architectural resilience to generative dilution comparing DenseNet-121 and EfficientNet-B3."
    )
    
    add_p(
        "The declining accuracy trajectories reflect this systematic degradation. DenseNet-121 suffered a substantial performance drop "
        "in test accuracy from 95.67% to 93.22%. EfficientNet-B3 exhibited greater robustness against generative noise propagation, but "
        "still suffered an absolute decline of 0.94%, falling from 96.98% to 96.04%."
    )
    
    add_figure(
        os.path.join(IMG_DIR, "p6_img44.jpeg"),
        "Fig. 12. Experiment 4: Per-class F1 degradation at maximum synthetic saturation across both backbones."
    )
    
    add_p(
        "Granular class-level heatmaps demonstrate that dilution was widespread, systematically degrading precision and recall across "
        "all diagnostic categories. Ultimately, these results establish a strict operational boundary for clinical AI: synthetic "
        "histopathology effectively regularizes minority imbalances, but overwhelming authentic tissue distributions actively impairs "
        "diagnostic performance."
    )
    
    # 7. Section VI
    add_heading_1("VI. Discussion: Pathological Failure Modes")
    add_p(
        "To investigate the exact failure mechanics of the pure-synthetic regime, we examined the normalized confusion matrices and "
        "cross-class predictions from Experiment 2."
    )
    
    add_figure(
        os.path.join(IMG_DIR, "p6_img45.png"),
        "Fig. 13. DenseNet-121: Normalized confusion matrix under pure-synthetic training, highlighting severe cross-class collapse into the Normal category."
    )
    add_figure(
        os.path.join(IMG_DIR, "p6_img46.png"),
        "Fig. 14. EfficientNet-B3: Normalized confusion matrix illustrating pervasive collapse of OSMF into normal tissue and PDOSCC into WDOSCC."
    )
    
    add_p(
        "The most severe failure occurred in the Oral Submucous Fibrosis (OSMF) class across both deep backbones. DenseNet-121 and "
        "EfficientNet-B3 erroneously classified 86.6% and 94.2% of authentic biopsies into the normal category, thereby collapsing the "
        "decision margin between healthy tissue and premalignant mucosal transformation. While the generator mimicked the general "
        "eosinophilic appearance of connective tissue, it failed to resolve the dense, bundled collagen fibrillogenesis and progressive "
        "stromal hyalinization that pathognomonically demarcate OSMF from normal mucosal architecture. A secondary diagnostic breakdown "
        "emerged within the OSCC grading hierarchy. EfficientNet-B3 misclassified 74.9% of aggressive PDOSCC cases as WDOSCC, indicating "
        "a collapse in tumour-grade separation. Clinical grading is determined by precise cellular features, especially nuclear-to-cytoplasmic "
        "ratios and the identification of distinct keratin pearls. The generated samples replaced discrete cellular phenotypes with smoothed "
        "malignant textures, leaving the classifiers unable to detect genuine diagnostic markers."
    )
    
    # 8. Section VII
    add_heading_1("VII. Conclusion")
    add_p(
        "This research reveals a critical vulnerability dissociation in generative medical AI: visual realism is clinically deceptive "
        "and masks the loss of key diagnostic features. Although StyleGAN2-ADA achieved remarkable qualitative fidelity and a low "
        "Fréchet Inception Distance (FID = 6.98), the synthetic data failed pure-synthetic stress tests, lacking the microscopic "
        "biological variance needed for sustaining independent clinical-grade classification. However, synthetic data remains highly "
        "valuable when strictly constrained: our empirical mixing experiments demonstrate that bounded augmentation functions as an "
        "effective regularizer to resolve minority-class imbalances. In contrast, aggressive synthetic over-saturation dilutes genuine "
        "clinical signals, compromising the classifiers’ discriminative margins. In the final analysis, high generative fidelity cannot "
        "be conflated with clinical validity. Future research must look beyond natural-image FID scores and necessitate downstream clinical "
        "stress tests before deploying synthetic data prior to clinical deployment."
    )
    
    # 9. References
    add_heading_1("References")
    references = [
        "[1] G. Huang, Z. Liu, L. Van Der Maaten, and K. Q. Weinberger, \"Densely connected convolutional networks,\" in Proc. IEEE Conf. Comput. Vis. Pattern Recognit. (CVPR), 2017, pp. 4700-4708.",
        "[2] M. Tan and Q. V. Le, \"EfficientNet: Rethinking model scaling for convolutional neural networks,\" in Int. Conf. Mach. Learn. (ICML), 2019, pp. 6105-6114.",
        "[3] L. Oakden-Rayner et al., \"Hidden stratification causes clinically meaningful failures in machine learning for medical imaging,\" in Proc. ACM Conf. Health, Inference, and Learning, 2020, pp. 151-159.",
        "[4] R. Chaudhary et al., \"ORCHID: A multi-center open-access oral cancer histology image database,\" Sci. Data, vol. 11, p. 123, 2024.",
        "[5] S. Chakrabarty et al., \"YOLO26/CNN benchmarking on the ORCHID oral cancer dataset,\" Preprint, 2026.",
        "[6] R. Chaudhary et al., \"OralPatho framework on whole slide images,\" medRxiv, 2023.",
        "[7] A. Kumar, B. Kumar, and C. Jindal, \"EfficientNet-B0 on balanced 3k ORCHID subset,\" IEEE Trans. Med. Imaging, 2026.",
        "[8] Y. Ren, X. Li, and Z. Chen, \"Computer-aided classification study for OSCC,\" Front. Oncol., vol. 15, 2026.",
        "[9] A. Tasnim et al., \"Histo-AdaptiveViT: Long-tailed classification for histopathology,\" in Proc. IEEE Int. Symp. Biomed. Imaging (ISBI), 2026.",
        "[10] I. Goodfellow et al., \"Generative adversarial nets,\" Adv. Neural Inf. Process. Syst. (NeurIPS), vol. 27, 2014.",
        "[11] T. Karras et al., \"Training generative adversarial networks with limited data,\" Adv. Neural Inf. Process. Syst. (NeurIPS), vol. 33, pp. 12104-12114, 2020.",
        "[12] J. Dee et al., \"Domain adaptation with StyleGAN2 on thyroid histopathology,\" PLOS ONE, vol. 19, 2024.",
        "[13] M. Heusel et al., \"GANs trained by a two time-scale update rule converge to a local Nash equilibrium,\" Adv. Neural Inf. Process. Syst., vol. 30, 2017.",
        "[14] M. Chong and D. Forsyth, \"Effectively unbiased FID and inception score and where to find them,\" in Proc. IEEE/CVF Conf. Comput. Vis. Pattern Recognit. (CVPR), 2020, pp. 6070-6079.",
        "[15] Y. Xue et al., \"Selective synthetic augmentation with HistoGAN for breast cancer pathology,\" Med. Image Anal., vol. 67, 2021.",
        "[16] M. Frid-Adar et al., \"GAN-based synthetic medical image augmentation for increased CNN performance in liver lesion classification,\" Neurocomputing, vol. 321, pp. 321-331, 2018."
    ]
    
    for ref in references:
        pref = doc.add_paragraph()
        pref.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pref.paragraph_format.space_before = Pt(0)
        pref.paragraph_format.space_after = Pt(2)
        pref.paragraph_format.line_spacing = 1.0
        pref.paragraph_format.left_indent = Inches(0.25)
        pref.paragraph_format.first_line_indent = Inches(-0.25)
        
        rr = pref.add_run(ref)
        rr.font.name = 'Times New Roman'
        rr.font.size = Pt(8)

    docx_path = os.path.join(BASE_DIR, 'IEEE_Paper.docx')
    doc.save(docx_path)
    # Also save as IEEE_Format_Paper.docx
    doc.save(os.path.join(BASE_DIR, 'IEEE_Format_Paper.docx'))
    print(f"Docx saved to: {docx_path}")

if __name__ == '__main__':
    build_docx()
